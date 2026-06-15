#-*- coding: utf-8 -*-
from odoo import models, fields, api, _
from odoo.exceptions import UserError
import base64
import openpyxl
from io import BytesIO

class UpdateReplenishmentWizard(models.TransientModel):
    _name = 'update.replenishment.wizard'
    _description = 'Update Replenishment Rules Wizard'

    company_ids = fields.Many2many('res.company', string='Compañías', required=True)
    file = fields.Binary(string='Excel', required=True)
    filename = fields.Char(string='Nombre del archivo')

    def action_update_replenishment(self):
        if not self.file:
            raise UserError("Please upload an Excel file.")

        # Read the uploaded Excel file with openpyxl
        file_data = base64.b64decode(self.file)
        wb = openpyxl.load_workbook(BytesIO(file_data), data_only=True)
        ws = wb.active

        # Read header row to find column positions
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        if 'CODIGO DE BARRA' not in headers or 'PRECIO USD' not in headers:
            raise UserError("The Excel file must contain 'CODIGO DE BARRA' and 'PRECIO USD' columns.")

        col_barcode = headers.index('CODIGO DE BARRA')
        col_price = headers.index('PRECIO USD')

        # Iterate over data rows (skip header)
        for row in ws.iter_rows(min_row=2, values_only=True):
            barcode_val = row[col_barcode]
            price_usd = row[col_price]
            if barcode_val is None:
                continue
            barcode = str(int(barcode_val)).strip()

            # Search for the replenishment rule for the given product and companies
            replenishment_rule = self.env['stock.warehouse.orderpoint'].search([
                ('product_id.barcode', '=', barcode),
                ('company_id', 'in', self.company_ids.ids)
            ], limit=1)

            if replenishment_rule:
                replenishment_rule.purchase_price = price_usd
            else:
                raise UserError(f"No replenishment rule found for barcode: {barcode}")

        action = {
            'name': _('Smart Pro'),
            'view_mode': 'tree,form',
            'res_model': 'stock.warehouse.orderpoint',
            'type': 'ir.actions.act_window',
            "domain": [],
            'target': 'current',
        }
        return action

